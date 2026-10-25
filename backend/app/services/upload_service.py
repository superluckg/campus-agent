"""上传边界校验、落盘与文本抽取。"""

from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from fastapi import UploadFile
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

from app.db.repositories.uploads import UploadRepository
from app.tools.document_reader import DocumentReadRequest, DocumentReaderService

from .upload_resolver import ScopedUploadResolver


class UploadService:
    def __init__(
        self,
        upload_root: Path,
        repository: UploadRepository,
        allowed_extensions: list[str],
        max_files: int,
        max_file_size_mb: int,
        max_output_characters: int,
        max_pdf_pages: int = 20,
        max_pptx_slides: int = 20,
        max_excel_sheets: int = 5,
        max_excel_rows_per_sheet: int = 500,
    ) -> None:
        self.upload_root = upload_root.resolve()
        self.repository = repository
        self.allowed_extensions = {".%s" % value.casefold().lstrip(".") for value in allowed_extensions}
        self.max_files = max_files
        self.max_bytes = max_file_size_mb * 1024 * 1024
        self.max_output_characters = max_output_characters
        self.max_pdf_pages = max_pdf_pages
        self.max_pptx_slides = max_pptx_slides
        self.max_excel_sheets = max_excel_sheets
        self.max_excel_rows_per_sheet = max_excel_rows_per_sheet

    async def save(
        self,
        user_id: str,
        conversation_id: str,
        upload: UploadFile,
    ):
        if await self.repository.count(user_id, conversation_id) >= self.max_files:
            raise ValueError("当前对话上传文件数量已达到上限。")
        original_name = Path(upload.filename or "unnamed").name
        suffix = Path(original_name).suffix.casefold()
        if suffix not in self.allowed_extensions:
            raise ValueError("不支持该文件格式。")
        content = await upload.read(self.max_bytes + 1)
        if len(content) > self.max_bytes:
            raise ValueError("文件大小超过限制。")
        if not content:
            raise ValueError("上传文件为空。")
        relative_dir = Path(conversation_id)
        stored_name = "%s%s" % (uuid4(), suffix)
        relative_path = relative_dir / stored_name
        destination = (self.upload_root / relative_path).resolve()
        destination.relative_to(self.upload_root)
        self._validate_structure(content, suffix)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(content)
        item = await self.repository.create(
            user_id=user_id,
            conversation_id=conversation_id,
            original_filename=original_name,
            relative_path=relative_path.as_posix(),
            extension=suffix.lstrip("."),
            mime_type=upload.content_type,
            size_bytes=len(content),
            content_hash=hashlib.sha256(content).hexdigest(),
        )
        resolver = ScopedUploadResolver(
            self.upload_root,
            [{"file_id": str(item.file_id), "relative_path": relative_path.as_posix()}],
        )
        try:
            result = DocumentReaderService(resolver, self.max_output_characters).read(
                DocumentReadRequest(file_ids=[str(item.file_id)], mode="full")
            )
            await self.repository.save_extracted_text(
                str(item.file_id), result.content, {"summaries": result.file_summaries}
            )
        except Exception as exc:
            await self.repository.mark_failed(str(item.file_id), str(exc))
            raise
        return item

    def _validate_structure(self, content: bytes, suffix: str) -> None:
        if suffix == ".pdf":
            pages = len(PdfReader(BytesIO(content)).pages)
            if pages > self.max_pdf_pages:
                raise ValueError("PDF页数超过%d页限制。" % self.max_pdf_pages)
        elif suffix == ".pptx":
            slides = len(Presentation(BytesIO(content)).slides)
            if slides > self.max_pptx_slides:
                raise ValueError("PPTX页数超过%d页限制。" % self.max_pptx_slides)
        elif suffix == ".xlsx":
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                if len(workbook.sheetnames) > self.max_excel_sheets:
                    raise ValueError("Excel工作表数量超过限制。")
                if any(workbook[name].max_row > self.max_excel_rows_per_sheet for name in workbook.sheetnames):
                    raise ValueError("Excel单个工作表行数超过限制。")
            finally:
                workbook.close()
