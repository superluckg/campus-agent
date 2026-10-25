"""把公共资料解析成带页码/幻灯片号的纯文本段。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader
from pptx import Presentation


@dataclass(frozen=True)
class ParsedSection:
    text: str
    page_number: int = None
    section_path: str = None


def parse_public_document(path: Path) -> list[ParsedSection]:
    suffix = path.suffix.casefold()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        sections = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append(ParsedSection(text=text, page_number=index))
        if not sections:
            raise ValueError("PDF没有可提取文本，可能是扫描版。")
        return sections
    if suffix == ".docx":
        document = DocxDocument(str(path))
        sections = []
        current_heading = None
        buffer = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style.name.startswith("Heading"):
                if buffer:
                    sections.append(
                        ParsedSection(text="\n".join(buffer), section_path=current_heading)
                    )
                    buffer = []
                current_heading = text
            else:
                buffer.append(text)
        if buffer:
            sections.append(ParsedSection(text="\n".join(buffer), section_path=current_heading))
        return sections
    if suffix == ".pptx":
        presentation = Presentation(str(path))
        return [
            ParsedSection(
                text="\n".join(
                    shape.text.strip()
                    for shape in slide.shapes
                    if hasattr(shape, "text") and shape.text.strip()
                ),
                page_number=index,
                section_path=(slide.shapes.title.text.strip() if slide.shapes.title else None),
            )
            for index, slide in enumerate(presentation.slides, start=1)
            if any(hasattr(shape, "text") and shape.text.strip() for shape in slide.shapes)
        ]
    if suffix in {".md", ".txt"}:
        return [ParsedSection(text=path.read_text(encoding="utf-8-sig"))]
    raise ValueError("公共知识库不支持格式：%s" % suffix)
