"""读取会话上传文件；不读取URL、不执行OCR、不进入公共知识库。"""

from __future__ import annotations

import base64
import json
from pathlib import Path
from typing import Any, Dict, List, Optional, Protocol

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field
from pypdf import PdfReader
from pptx import Presentation


class UploadResolver(Protocol):
    def resolve(self, file_id: str) -> Path: ...


class DocumentReadRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    file_ids: List[str] = Field(min_length=1, max_length=10)
    mode: str
    page_ranges: Dict[str, List[int]] = Field(default_factory=dict)
    sheet_names: Dict[str, List[str]] = Field(default_factory=dict)
    query: Optional[str] = None
    continuation_token: Optional[str] = None


class DocumentReadSegment(BaseModel):
    content: str
    metadata: Dict[str, Any]


class DocumentReadResult(BaseModel):
    model_config = ConfigDict(extra="forbid")

    content: str
    file_summaries: List[Dict[str, Any]] = Field(default_factory=list)
    truncated: bool = False
    continuation_token: Optional[str] = None
    segments: List[DocumentReadSegment] = Field(default_factory=list)


class DocumentReaderService:
    def __init__(self, resolver: UploadResolver, max_output_characters: int = 40000) -> None:
        self.resolver = resolver
        self.max_output_characters = max_output_characters

    def read(self, request: DocumentReadRequest) -> DocumentReadResult:
        if request.mode not in {"outline", "pages", "full"}:
            raise ValueError("mode必须是outline、pages或full。")
        blocks: list[str] = []
        summaries: list[Dict[str, Any]] = []
        provenance: list[Dict[str, Any]] = []
        for file_id in request.file_ids:
            path = self.resolver.resolve(file_id)
            parts, summary = self._extract(path, request, file_id)
            summary["file_id"] = file_id
            summaries.append(summary)
            blocks.extend(parts)
            numbers = (request.page_ranges.get(file_id) or list(range(1, len(parts) + 1)))
            for index, part in enumerate(parts):
                metadata = {"file_id": file_id, "filename": path.name,
                            "title": path.name, "source_type": "conversation_upload"}
                if request.mode != "outline" and path.suffix.casefold() in {".pdf", ".pptx"}:
                    metadata["page_number"] = numbers[index]
                if path.suffix.casefold() == ".xlsx" and request.mode != "outline":
                    metadata["section_path"] = summary["sheets"][index]
                provenance.append(metadata)
        full_text = "\n\n".join(blocks)
        offset = self._decode_token(request.continuation_token)
        remaining = full_text[offset:]
        content = remaining[: self.max_output_characters]
        truncated = len(remaining) > self.max_output_characters
        token = self._encode_token(offset + len(content)) if truncated else None
        segments = []
        position = 0
        for block, metadata in zip(blocks, provenance, strict=True):
            start, end = max(offset, position), min(offset + len(content), position + len(block))
            if start < end:
                segments.append(DocumentReadSegment(content=full_text[start:end], metadata={
                    **metadata, "slice_start": start - position, "slice_end": end - position,
                    "read_mode": request.mode,
                }))
            position += len(block) + 2
        return DocumentReadResult(
            content=content,
            file_summaries=summaries,
            truncated=truncated,
            continuation_token=token,
            segments=segments,
        )

    def _extract(
        self, path: Path, request: DocumentReadRequest, file_id: str
    ) -> tuple[list[str], Dict[str, Any]]:
        suffix = path.suffix.casefold()
        if suffix == ".pdf":
            return self._pdf(path, request, file_id)
        if suffix == ".docx":
            return self._docx(path, request)
        if suffix == ".pptx":
            return self._pptx(path, request, file_id)
        if suffix in {".md", ".txt"}:
            return self._text(path, request)
        if suffix == ".xlsx":
            return self._xlsx(path, request, file_id)
        raise ValueError("不支持的文件格式：%s" % suffix)

    @staticmethod
    def _selected_numbers(
        configured: Dict[str, List[int]], file_id: str, total: int
    ) -> list[int]:
        return configured.get(file_id) or list(range(1, total + 1))

    def _pdf(self, path: Path, request: DocumentReadRequest, file_id: str):
        reader = PdfReader(str(path))
        pages = self._selected_numbers(request.page_ranges, file_id, len(reader.pages))
        if request.mode == "outline":
            return ["%s：PDF，共%d页。" % (path.name, len(reader.pages))], {
                "filename": path.name,
                "pages": len(reader.pages),
            }
        blocks = []
        for page_number in pages:
            if page_number < 1 or page_number > len(reader.pages):
                raise ValueError("%s不存在第%d页。" % (path.name, page_number))
            text = (reader.pages[page_number - 1].extract_text() or "").strip()
            if not text:
                raise ValueError("%s第%d页无法提取文本，可能是扫描版PDF。" % (path.name, page_number))
            blocks.append("[%s 第%d页]\n%s" % (path.name, page_number, text))
        return blocks, {"filename": path.name, "pages": len(reader.pages)}

    @staticmethod
    def _docx(path: Path, request: DocumentReadRequest):
        document = DocxDocument(str(path))
        headings = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.style.name.startswith("Heading") and paragraph.text.strip()
        ]
        if request.mode == "outline":
            text = "%s：DOCX；标题：%s" % (
                path.name,
                "、".join(headings) or "未识别到标题",
            )
            return [text], {"filename": path.name, "headings": headings}
        content = "\n".join(
            paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()
        )
        return ["[%s]\n%s" % (path.name, content)], {
            "filename": path.name,
            "headings": headings,
        }

    def _pptx(self, path: Path, request: DocumentReadRequest, file_id: str):
        presentation = Presentation(str(path))
        slides = self._selected_numbers(
            request.page_ranges, file_id, len(presentation.slides)
        )
        if request.mode == "outline":
            titles = [
                slide.shapes.title.text.strip() if slide.shapes.title else "未命名"
                for slide in presentation.slides
            ]
            return [
                "%s：PPTX，共%d页；%s"
                % (path.name, len(presentation.slides), "、".join(titles))
            ], {"filename": path.name, "slides": len(presentation.slides)}
        blocks = []
        for slide_number in slides:
            if slide_number < 1 or slide_number > len(presentation.slides):
                raise ValueError("%s不存在第%d页幻灯片。" % (path.name, slide_number))
            slide = presentation.slides[slide_number - 1]
            text = "\n".join(
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            )
            blocks.append("[%s 第%d页]\n%s" % (path.name, slide_number, text))
        return blocks, {"filename": path.name, "slides": len(presentation.slides)}

    @staticmethod
    def _text(path: Path, request: DocumentReadRequest):
        content = path.read_text(encoding="utf-8-sig")
        if request.mode == "outline":
            headings = [
                line.strip() for line in content.splitlines() if line.lstrip().startswith("#")
            ]
            return [
                "%s：文本文件；标题：%s"
                % (path.name, "、".join(headings[:20]) or "无")
            ], {"filename": path.name, "characters": len(content)}
        return ["[%s]\n%s" % (path.name, content)], {
            "filename": path.name,
            "characters": len(content),
        }

    @staticmethod
    def _xlsx(path: Path, request: DocumentReadRequest, file_id: str):
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            selected = request.sheet_names.get(file_id) or workbook.sheetnames
            if request.mode == "outline":
                return [
                    "%s：Excel；工作表：%s" % (path.name, "、".join(workbook.sheetnames))
                ], {"filename": path.name, "sheets": workbook.sheetnames}
            blocks = []
            for sheet_name in selected:
                if sheet_name not in workbook.sheetnames:
                    raise ValueError("%s不存在工作表%s。" % (path.name, sheet_name))
                sheet = workbook[sheet_name]
                lines = [
                    "\t".join("" if value is None else str(value) for value in row)
                    for row in sheet.iter_rows(values_only=True)
                ]
                blocks.append("[%s / %s]\n%s" % (path.name, sheet_name, "\n".join(lines)))
            return blocks, {"filename": path.name, "sheets": selected}
        finally:
            workbook.close()

    @staticmethod
    def _encode_token(offset: int) -> str:
        raw = json.dumps({"offset": offset}).encode("utf-8")
        return base64.urlsafe_b64encode(raw).decode("ascii")

    @staticmethod
    def _decode_token(token: Optional[str]) -> int:
        if not token:
            return 0
        try:
            payload = json.loads(base64.urlsafe_b64decode(token.encode("ascii")))
            offset = int(payload["offset"])
            if offset < 0:
                raise ValueError
            return offset
        except Exception as exc:
            raise ValueError("continuation_token无效。") from exc

