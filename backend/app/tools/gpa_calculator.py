"""确定性 GPA Calculator；不允许 LLM 参与计算。"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable, Mapping, Protocol

from openpyxl import load_workbook


REQUIRED_COLUMNS = ("课程名称", "学分", "成绩", "课程类别", "修读状态")
EXCLUDED_CATEGORY_TERMS = ("体育", "公共选修", "通识选修", "任选")
INCLUDED_CATEGORY_TERMS = ("主修", "专业", "必修")
EXCLUDED_STATUS_TERMS = ("重修", "缓考", "pass", "fail", "通过", "不通过")


class UploadPathResolver(Protocol):
    def resolve(self, file_id: str) -> Path: ...


@dataclass(frozen=True)
class CourseRow:
    name: str
    credits: Decimal
    score: Decimal
    category: str
    status: str


@dataclass(frozen=True)
class GPAResult:
    gpa: Decimal
    included_courses: int
    excluded_courses: int
    total_credits: Decimal

    def user_text(self) -> str:
        return "GPA：%s" % format(self.gpa, ".2f")


def score_to_point(score: Decimal) -> Decimal:
    if score >= 90:
        return Decimal("4.0")
    if score >= 85:
        return Decimal("3.7")
    if score >= 82:
        return Decimal("3.3")
    if score >= 78:
        return Decimal("3.0")
    if score >= 75:
        return Decimal("2.7")
    if score >= 72:
        return Decimal("2.3")
    if score >= 68:
        return Decimal("2.0")
    if score >= 64:
        return Decimal("1.5")
    if score >= 60:
        return Decimal("1.0")
    return Decimal("0.0")


def _is_included(row: CourseRow) -> bool:
    category = row.category.casefold()
    status = row.status.casefold()
    if any(term.casefold() in category for term in EXCLUDED_CATEGORY_TERMS):
        return False
    if any(term.casefold() in status for term in EXCLUDED_STATUS_TERMS):
        return False
    return any(term.casefold() in category for term in INCLUDED_CATEGORY_TERMS)


def calculate_gpa(rows: Iterable[Mapping[str, object]]) -> GPAResult:
    included: list[CourseRow] = []
    excluded = 0
    for row_number, raw in enumerate(rows, start=2):
        missing = [name for name in REQUIRED_COLUMNS if name not in raw]
        if missing:
            raise ValueError("成绩单缺少列：%s" % "、".join(missing))
        try:
            credits = Decimal(str(raw["学分"]))
            score = Decimal(str(raw["成绩"]))
        except Exception as exc:
            raise ValueError("第%d行的学分或成绩不是数值。" % row_number) from exc
        if credits <= 0:
            raise ValueError("第%d行的学分必须大于0。" % row_number)
        if score < 0 or score > 100:
            raise ValueError("第%d行的成绩必须处于0到100之间。" % row_number)
        course = CourseRow(
            name=str(raw["课程名称"]).strip(),
            credits=credits,
            score=score,
            category=str(raw["课程类别"]).strip(),
            status=str(raw["修读状态"]).strip(),
        )
        if _is_included(course):
            included.append(course)
        else:
            excluded += 1
    if not included:
        raise ValueError("成绩单中没有符合演示规则的主修课程。")
    total_credits = sum((course.credits for course in included), Decimal("0"))
    weighted = sum(
        (score_to_point(course.score) * course.credits for course in included),
        Decimal("0"),
    )
    gpa = (weighted / total_credits).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return GPAResult(
        gpa=gpa,
        included_courses=len(included),
        excluded_courses=excluded,
        total_credits=total_credits,
    )


class GPACalculatorService:
    def __init__(self, resolver: UploadPathResolver) -> None:
        self.resolver = resolver

    def calculate(self, file_id: str) -> GPAResult:
        path = self.resolver.resolve(file_id)
        if path.suffix.casefold() != ".xlsx":
            raise ValueError("GPA Calculator只接受XLSX成绩单。")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = [str(value).strip() if value is not None else "" for value in next(rows)]
            except StopIteration as exc:
                raise ValueError("成绩单为空。") from exc
            missing = [name for name in REQUIRED_COLUMNS if name not in headers]
            if missing:
                raise ValueError("成绩单缺少列：%s" % "、".join(missing))
            records = [dict(zip(headers, values)) for values in rows if any(value is not None for value in values)]
            return calculate_gpa(records)
        finally:
            workbook.close()

